"""
converter.py — 多引擎文档转换器

支持8种引擎自动fallback:
  markitdown, docling, pymupdf, antiword, textutil, libreoffice, openpyxl, python-pptx

用法:
    from markitdown_enhanced import DocumentConverter
    converter = DocumentConverter()
    md_text = converter.convert("input.docx")
"""

import logging
import os
import subprocess
import tempfile
import time
from pathlib import Path
from typing import Optional

log = logging.getLogger("markitdown_enhanced.converter")

# 配置HuggingFace镜像（国内环境）
os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")

# 文件大小阈值: 超过此值优先使用PyMuPDF（更快）
_LARGE_FILE_THRESHOLD = 10 * 1024 * 1024  # 10MB


class DocumentConverter:
    """多引擎文档转换器，自动选择最佳引擎"""

    def __init__(self, preferred_engine: str = "auto"):
        """
        Args:
            preferred_engine: 首选引擎
                "auto" — 按文件类型和大小自动选择
                "markitdown" — 只用MarkItDown原生
                "docling" — 优先Docling
                "pymupdf" — 优先PyMuPDF
        """
        self.preferred_engine = preferred_engine

    def convert(self, file_path: str) -> str:
        """
        转换文档为Markdown文本。
        按优先级尝试多个引擎，第一个成功的就返回。

        引擎选择逻辑:
          PDF (<=10MB): markitdown → docling → pymupdf
          PDF (>10MB):  pymupdf → markitdown → docling
          DOCX:         markitdown → docling → libreoffice转码→markitdown
          DOC:          antiword → textutil → libreoffice转码→docling
          XLSX:         markitdown → openpyxl手动提取
          PPTX:         markitdown → python-pptx手动提取
          其他:         markitdown → docling

        Returns:
            转换后的Markdown文本
        Raises:
            RuntimeError: 所有引擎都失败时
        """
        file_path = os.path.abspath(file_path)
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = Path(file_path).suffix.lower()
        file_size = os.path.getsize(file_path)

        log.info("开始转换: %s (类型: %s, 大小: %.1fMB)",
                 os.path.basename(file_path), ext, file_size / 1024 / 1024)

        if self.preferred_engine == "markitdown":
            # 只用原生引擎
            engines = [self._try_markitdown]
        elif self.preferred_engine == "docling":
            engines = [self._try_docling, self._try_markitdown, self._try_pymupdf]
        elif self.preferred_engine == "pymupdf":
            engines = [self._try_pymupdf, self._try_markitdown, self._try_docling]
        elif ext == ".pdf":
            if file_size > _LARGE_FILE_THRESHOLD:
                engines = [self._try_pymupdf, self._try_markitdown, self._try_docling]
            else:
                engines = [self._try_markitdown, self._try_docling, self._try_pymupdf]
        elif ext == ".docx":
            engines = [self._try_markitdown, self._try_docling, self._try_libreoffice_docx]
        elif ext == ".doc":
            engines = [self._try_antiword, self._try_textutil, self._try_libreoffice_doc]
        elif ext in (".xlsx", ".xls"):
            engines = [self._try_markitdown, self._try_openpyxl]
        elif ext == ".pptx":
            engines = [self._try_markitdown, self._try_python_pptx]
        else:
            engines = [self._try_markitdown, self._try_docling]

        last_error = None
        for engine_fn in engines:
            engine_name = engine_fn.__name__.replace("_try_", "")
            log.info("尝试引擎: %s", engine_name)
            try:
                result = engine_fn(file_path)
                if result and len(result.strip()) >= 10:
                    log.info("引擎 %s 成功: %d字符", engine_name, len(result))
                    return result
                else:
                    log.warning("引擎 %s 返回内容过短(%d字符), 跳过",
                                engine_name, len(result.strip()) if result else 0)
            except Exception as e:
                log.warning("引擎 %s 失败: %s", engine_name, e)
                last_error = e

        raise RuntimeError(
            f"所有引擎均失败: {file_path}\n"
            f"尝试了: {[fn.__name__.replace('_try_', '') for fn in engines]}\n"
            f"最后错误: {last_error}"
        )

    # =========================================================================
    # 引擎方法
    # =========================================================================

    def _try_markitdown(self, file_path: str) -> Optional[str]:
        """用MarkItDown原生引擎"""
        from markitdown import MarkItDown
        md = MarkItDown()
        result = md.convert(file_path)
        text = result.text_content or ""
        if len(text.strip()) < 10:
            return None
        return text

    def _try_docling(self, file_path: str) -> Optional[str]:
        """用Docling引擎（需要docling包）"""
        try:
            from docling.document_converter import DocumentConverter
            converter = DocumentConverter()
            result = converter.convert(file_path)
            md = result.document.export_to_markdown()
            if len(md.strip()) < 10:
                return None
            return md
        except ImportError:
            log.debug("docling未安装")
            return None
        except Exception as e:
            log.debug("docling转换异常: %s", e)
            return None

    def _try_pymupdf(self, file_path: str) -> Optional[str]:
        """用PyMuPDF提取PDF文本"""
        try:
            import fitz
            doc = fitz.open(file_path)
            pages = []
            for page in doc:
                text = page.get_text()
                if text.strip():
                    pages.append(text)
            doc.close()
            md = "\n\n".join(pages)
            if len(md.strip()) < 10:
                return None
            return md
        except ImportError:
            log.debug("fitz(PyMuPDF)未安装")
            return None
        except Exception as e:
            log.debug("pymupdf提取异常: %s", e)
            return None

    def _try_antiword(self, file_path: str) -> Optional[str]:
        """用antiword提取.doc文件"""
        try:
            result = subprocess.run(
                ["antiword", file_path],
                capture_output=True, text=True, timeout=30,
            )
            if result.returncode != 0:
                return None
            text = result.stdout
            if len(text.strip()) < 10:
                return None
            return text
        except (FileNotFoundError, subprocess.TimeoutExpired):
            return None
        except Exception as e:
            log.debug("antiword异常: %s", e)
            return None

    def _try_textutil(self, file_path: str) -> Optional[str]:
        """用macOS textutil提取（.doc/.docx）"""
        try:
            with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
                out_path = f.name
            try:
                subprocess.run(
                    ["textutil", "-convert", "txt", "-output", out_path, file_path],
                    capture_output=True, timeout=30,
                )
                with open(out_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
                if len(text.strip()) < 10:
                    return None
                return text
            finally:
                if os.path.exists(out_path):
                    os.unlink(out_path)
        except Exception as e:
            log.debug("textutil异常: %s", e)
            return None

    def _try_libreoffice_docx(self, file_path: str) -> Optional[str]:
        """用LibreOffice将DOCX重新转码后再提取"""
        try:
            with tempfile.TemporaryDirectory(prefix="lo_convert_") as tmpdir:
                subprocess.run(
                    ["/opt/homebrew/bin/soffice", "--headless", "--convert-to", "docx",
                     "--outdir", tmpdir, file_path],
                    capture_output=True, timeout=120,
                )
                docx_files = [f for f in os.listdir(tmpdir) if f.endswith(".docx")]
                if not docx_files:
                    return None
                converted = os.path.join(tmpdir, docx_files[0])
                # 用markitdown提取转换后的文件
                text = self._try_markitdown(converted)
                if not text:
                    text = self._try_docling(converted)
                return text
        except Exception as e:
            log.debug("libreoffice(docx)异常: %s", e)
            return None

    def _try_libreoffice_doc(self, file_path: str) -> Optional[str]:
        """用LibreOffice将.doc转码后再提取"""
        try:
            with tempfile.TemporaryDirectory(prefix="lo_convert_") as tmpdir:
                subprocess.run(
                    ["/opt/homebrew/bin/soffice", "--headless", "--convert-to", "docx",
                     "--outdir", tmpdir, file_path],
                    capture_output=True, timeout=120,
                )
                docx_files = [f for f in os.listdir(tmpdir) if f.endswith(".docx")]
                if not docx_files:
                    return None
                converted = os.path.join(tmpdir, docx_files[0])
                text = self._try_docling(converted)
                if not text:
                    text = self._try_markitdown(converted)
                return text
        except Exception as e:
            log.debug("libreoffice(doc)异常: %s", e)
            return None

    def _try_openpyxl(self, file_path: str) -> Optional[str]:
        """用openpyxl手动提取xlsx"""
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True, data_only=True)
            parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                parts.append(f"## {sheet_name}\n")
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    header = [str(c) if c is not None else "" for c in rows[0]]
                    parts.append("| " + " | ".join(header) + " |")
                    parts.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for row in rows[1:101]:
                        cells = [str(c) if c is not None else "—" for c in row]
                        parts.append("| " + " | ".join(cells) + " |")
                parts.append("")
            wb.close()
            md = "\n".join(parts)
            if len(md.strip()) < 10:
                return None
            return md
        except Exception as e:
            log.debug("openpyxl异常: %s", e)
            return None

    def _try_python_pptx(self, file_path: str) -> Optional[str]:
        """用python-pptx手动提取pptx"""
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                parts.append(f"### Slide {i}\n")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = para.text.strip()
                            if text:
                                parts.append(text)
                parts.append("")
            md = "\n".join(parts)
            if len(md.strip()) < 10:
                return None
            return md
        except Exception as e:
            log.debug("python-pptx异常: %s", e)
            return None


def extract_docx_images(file_path: str) -> list:
    """
    从DOCX文件中直接提取嵌入的PNG/JPEG图片。

    Args:
        file_path: DOCX文件路径

    Returns:
        按文档顺序排列的图片列表:
        [{"mime_type": str, "data": bytes, "alt_text": str}, ...]
        仅返回PNG/JPEG格式，跳过EMF/WMF等不可识别格式。
    """
    try:
        from docx import Document
    except ImportError:
        log.debug("python-docx未安装，跳过DOCX图片提取")
        return []

    images = []
    try:
        doc = Document(file_path)

        # 按文档顺序遍历paragraph中的inline shapes
        for para in doc.paragraphs:
            for run in para.runs:
                if run._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
                    # 这个run包含图片，查找对应的relationship
                    drawings = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                    for blip in drawings:
                        embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                        if embed_id and embed_id in doc.part.rels:
                            rel = doc.part.rels[embed_id]
                            ct = rel.target_part.content_type
                            if ct in ("image/png", "image/jpeg", "image/jpg"):
                                blob = rel.target_part.blob
                                # 获取alt text
                                alt = ""
                                desc_elems = blip.getparent().findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')
                                if desc_elems:
                                    alt = desc_elems[0].get('descr', '') or desc_elems[0].get('name', '')
                                images.append({
                                    "mime_type": ct,
                                    "data": blob,
                                    "alt_text": alt,
                                })
    except Exception as e:
        log.warning("DOCX图片提取异常: %s", e)

    log.info("DOCX图片提取: 共%d张PNG/JPEG (跳过EMF/WMF)", len(images))
    return images


# ============================================================================
# PPTX 视觉提取管线
# ============================================================================

_PPTX_VISION_PROMPT = """请详细描述这页PPT的内容，用简洁的结构化文本输出。

要求：
1. 先写页面标题（如有）
2. 列出所有可见文字内容（包括标注、注释、数据标签）
3. 描述图表类型（柱状图/饼图/流程图/架构图/表格等）及关键数据点
4. 如有截图，描述截图中的关键信息
5. 保留所有数字、百分比、金额等量化数据

直接输出自然语言描述，不要用JSON格式。"""


def extract_pptx_vision(
    pptx_path: str,
    describe_fn=None,
    dpi: int = 150,
    max_workers: int = 4,
) -> str:
    """
    PPTX → PDF → 逐页图片 → 视觉LLM → 合并描述

    比python-pptx纯文字提取更完整，能识别图表、架构图、截图等视觉内容。

    Args:
        pptx_path: PPTX文件路径
        describe_fn: 视觉描述函数，签名 describe_fn(base64_str, mime_type) -> str
        dpi: 渲染DPI，150足够清晰且传输快
        max_workers: 并发识别线程数

    Returns:
        合并后的逐页描述Markdown文本
    """
    import tempfile
    import subprocess
    import base64 as b64mod
    import concurrent.futures
    import fitz  # PyMuPDF

    LIBREOFFICE = "/opt/homebrew/bin/soffice"

    # 1. PPTX → PDF
    log.info("[视觉管线] PPTX→PDF: %s", os.path.basename(pptx_path))
    tmp_dir = tempfile.mkdtemp(prefix="pptx_vision_")

    try:
        result = subprocess.run(
            [LIBREOFFICE, "--headless", "--convert-to", "pdf", "--outdir", tmp_dir, pptx_path],
            capture_output=True, text=True, timeout=120,
        )
        if result.returncode != 0:
            raise RuntimeError(f"LibreOffice失败: {result.stderr[:200]}")

        pdf_name = Path(pptx_path).stem + ".pdf"
        pdf_path = os.path.join(tmp_dir, pdf_name)
        if not os.path.isfile(pdf_path):
            raise RuntimeError(f"PDF未生成: {pdf_path}")

        # 2. PDF → 逐页PNG
        log.info("[视觉管线] PDF→图片...")
        doc = fitz.open(pdf_path)
        zoom = dpi / 72
        mat = fitz.Matrix(zoom, zoom)

        image_paths = []
        for i in range(len(doc)):
            page = doc[i]
            pix = page.get_pixmap(matrix=mat)
            img_path = os.path.join(tmp_dir, f"page_{i+1:03d}.png")
            pix.save(img_path)
            image_paths.append(img_path)
        doc.close()
        log.info("[视觉管线] %d页图片", len(image_paths))

        # 3. 逐页视觉识别
        total = len(image_paths)
        page_results = [None] * total

        def _describe_page(page_idx):
            if not describe_fn:
                return f"[第{page_idx+1}页 - 无视觉识别函数]"
            with open(image_paths[page_idx], "rb") as f:
                img_b64 = b64mod.b64encode(f.read()).decode()
            try:
                # 调用视觉LLM（带prompt）
                from .harness.client import LLMSummarizer
                # describe_fn 是绑定方法，访问self获取client
                client = describe_fn.__self__.client if hasattr(describe_fn, '__self__') else None
                if client:
                    resp = client.chat.completions.create(
                        model=describe_fn.__self__.vision_model if hasattr(describe_fn.__self__, 'vision_model') else "glm-4v-flash",
                        messages=[{
                            "role": "user",
                            "content": [
                                {"type": "text", "text": _PPTX_VISION_PROMPT},
                                {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{img_b64}"}},
                            ]
                        }],
                        max_tokens=1000,
                        temperature=0.1,
                    )
                    return resp.choices[0].message.content.strip()
                else:
                    return describe_fn(img_b64, "image/png")
            except Exception as e:
                log.warning("[视觉管线] 第%d页识别失败: %s", page_idx + 1, e)
                return f"[第{page_idx+1}页识别失败]"

        log.info("[视觉管线] 开始识别%d页 (并发%d)...", total, max_workers)
        t0 = time.time()
        with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as ex:
            futures = {ex.submit(_describe_page, i): i for i in range(total)}
            for future in concurrent.futures.as_completed(futures):
                idx = futures[future]
                page_results[idx] = future.result()
                done = sum(1 for r in page_results if r is not None)
                if done % 10 == 0 or done == total:
                    log.info("[视觉管线] 进度 %d/%d", done, total)

        elapsed = time.time() - t0
        log.info("[视觉管线] 识别完成: %d页, %.1fs", total, elapsed)

        # 4. 合并为Markdown
        lines = []
        for i, desc in enumerate(page_results):
            lines.append(f"### 第{i+1}页\n")
            lines.append(desc if desc else "[空白页]")
            lines.append("")

        md_text = "\n".join(lines)
        log.info("[视觉管线] 输出 %d字符", len(md_text))
        return md_text

    finally:
        # 清理临时目录
        import shutil
        shutil.rmtree(tmp_dir, ignore_errors=True)
